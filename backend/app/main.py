from fastapi import FastAPI, Query, HTTPException, Body, UploadFile, File
from fastapi.responses import JSONResponse, StreamingResponse
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from typing import List, Optional
from datetime import datetime
from dotenv import load_dotenv
from io import BytesIO
import csv
import os

load_dotenv(os.path.join(os.path.dirname(__file__), '..', '.env'))

from . import repository as repo
from .api import saved_periods
from .services import catalog_service, price_service
from .services.catalog_service import DatabaseUnavailableError

app = FastAPI(title="鱼价数据平台 API", version="0.1.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# 注册保存时间段功能路由
app.include_router(saved_periods.router, prefix="/api", tags=["保存时间段"])

@app.get("/api/fishes")
def get_fishes(q: str | None = Query(default=None)):
    try:
        return catalog_service.list_fishes(q)
    except DatabaseUnavailableError as exc:
        raise HTTPException(status_code=503, detail=str(exc))

@app.get("/api/markets")
def get_markets(region_code: str | None = Query(default=None)):
    try:
        return catalog_service.list_markets(region_code)
    except DatabaseUnavailableError as exc:
        raise HTTPException(status_code=503, detail=str(exc))

@app.get("/api/prices")
def get_prices(
    fish_id: int, 
    market_id: int, 
    start: str | None = None, 
    end: str | None = None, 
    granularity: str = Query(default="day"),
    price_type: str = Query(default="pond"),
    unit: str | None = Query(default="kg"),
    lunar_mode: bool = Query(default=False),
    include_unverified: bool = Query(default=False)
):
    """
    查询价格数据
    
    正式查询只读取本地数据库，不触发外部采集或 mock 数据。
    默认只返回候选可信数据；如需查看待校验数据，显式传 include_unverified=true。
    """
    now = datetime.now().date().isoformat()
    if not start: 
        start = now
    if not end:
        end = now
    if granularity not in ("day","week","month"):
        granularity = "day"
    if price_type not in ("pond", "wholesale", "retail"):
        raise HTTPException(status_code=400, detail="price_type must be pond, wholesale, or retail")
    if unit not in ("kg", "斤", None):
        raise HTTPException(status_code=400, detail="unit must be kg or 斤")
    try:
        res = price_service.query_prices(
            fish_id=fish_id,
            market_id=market_id,
            start=start,
            end=end,
            granularity=granularity,
            price_type=price_type,
            unit=unit,
            lunar_mode=lunar_mode,
            include_unverified=include_unverified,
        )
    except DatabaseUnavailableError as exc:
        raise HTTPException(status_code=503, detail=str(exc))
    return JSONResponse(res)

@app.get("/api/forecast")
def get_forecast(fish_id: int, market_id: int, horizon: int = Query(default=14)):
    from . import data_loader as dl

    res = dl.forecast(fish_id, market_id, horizon)
    return JSONResponse(res)

class PriceInput(BaseModel):
    fish_id: int
    market_id: int
    price: float
    ts: str
    currency: str = "CNY"
    unit: str = "kg"
    price_type: str = "pond"
    source_url: Optional[str] = None
    source_text: Optional[str] = None

class BatchPriceInput(BaseModel):
    prices: List[PriceInput]

@app.post("/api/prices")
def add_price(data: PriceInput):
    try:
        datetime.fromisoformat(data.ts.replace('Z', '+00:00'))
    except:
        raise HTTPException(status_code=400, detail="Invalid date format. Use ISO format like 2024-01-15 or 2024-01-15T00:00:00")
    point = {
        "price": data.price,
        "ts": data.ts if 'T' in data.ts else data.ts + "T00:00:00+08:00",
        "currency": data.currency,
        "unit": data.unit
    }
    success = repo.insert_prices(data.fish_id, data.market_id, [point], data.price_type, data.source_url)
    if success:
        return {"success": True, "message": "Price added successfully"}
    raise HTTPException(status_code=500, detail="Failed to add price")

@app.post("/api/prices/batch")
def add_prices_batch(data: BatchPriceInput):
    added = 0
    failed = 0
    for p in data.prices:
        try:
            point = {
                "price": p.price,
                "ts": p.ts if 'T' in p.ts else p.ts + "T00:00:00+08:00",
                "currency": p.currency,
                "unit": p.unit
            }
            if repo.insert_prices(p.fish_id, p.market_id, [point], p.price_type, p.source_url):
                added += 1
            else:
                failed += 1
        except:
            failed += 1
    return {"success": True, "added": added, "failed": failed}

@app.get("/api/fishes/list")
def list_fishes():
    try:
        return catalog_service.list_fishes()
    except DatabaseUnavailableError as exc:
        raise HTTPException(status_code=503, detail=str(exc))

@app.get("/api/markets/list")
def list_markets():
    try:
        return catalog_service.list_markets()
    except DatabaseUnavailableError as exc:
        raise HTTPException(status_code=503, detail=str(exc))

@app.get("/api/templates/prices.xlsx")
def download_excel_template():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "价格数据"
    
    headers = ["鱼种名称", "市场名称", "价格", "日期", "货币", "单位", "价格类型", "来源URL"]
    ws.append(headers)
    
    example_data = [
        ["鲈鱼", "上海农产品中心", 15.80, "2024-01-15", "CNY", "kg", "pond", ""],
        ["泥鳅", "民众渔业", 12.50, "2024-01-16", "CNY", "kg", "pond", ""],
        ["草鱼", "上海农产品中心", 8.20, "2024-01-17", "CNY", "kg", "wholesale", ""]
    ]
    for row in example_data:
        ws.append(row)
    
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        ws.column_dimensions[column].width = max_length + 4
    
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=prices_template.xlsx"}
    )

@app.get("/api/templates/prices.csv")
def download_csv_template():
    csv_content = "鱼种名称,市场名称,价格,日期,货币,单位,价格类型,来源URL\n鲈鱼,上海农产品中心,15.80,2024-01-15,CNY,kg,pond,\n泥鳅,民众渔业,12.50,2024-01-16,CNY,kg,pond,\n草鱼,上海农产品中心,8.20,2024-01-17,CNY,kg,wholesale,"
    
    return StreamingResponse(
        BytesIO(csv_content.encode('utf-8-sig')),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": "attachment; filename=prices_template.csv"}
    )

def import_price_rows(rows):
    fishes = {f["name"]: f["id"] for f in list_fishes()}
    markets = {m["name"]: m["id"] for m in list_markets()}

    added = 0
    failed = 0
    errors = []

    for row_idx, row in rows:
        if not row or not row[0]:
            continue

        try:
            fish_name = str(row[0]).strip() if row[0] else None
            market_name = str(row[1]).strip() if row[1] else None
            price = float(row[2]) if row[2] else None
            ts = str(row[3]).strip() if row[3] else None

            if not all([fish_name, market_name, price, ts]):
                errors.append(f"第{row_idx}行: 缺少必填字段")
                failed += 1
                continue

            fish_id = fishes.get(fish_name)
            market_id = markets.get(market_name)

            if not fish_id:
                errors.append(f"第{row_idx}行: 找不到鱼种 '{fish_name}'")
                failed += 1
                continue
            if not market_id:
                errors.append(f"第{row_idx}行: 找不到市场 '{market_name}'")
                failed += 1
                continue

            price_type = str(row[6]).strip() if len(row) > 6 and row[6] else "pond"
            source_url = str(row[7]).strip() if len(row) > 7 and row[7] else None

            point = {
                "price": price,
                "ts": ts if 'T' in ts else ts + "T00:00:00+08:00",
                "currency": str(row[4]).strip() if len(row) > 4 and row[4] else "CNY",
                "unit": str(row[5]).strip() if len(row) > 5 and row[5] else "kg"
            }

            if repo.insert_prices(fish_id, market_id, [point], price_type, source_url):
                added += 1
            else:
                failed += 1
        except Exception as e:
            errors.append(f"第{row_idx}行: {str(e)}")
            failed += 1

    return {
        "success": True,
        "added": added,
        "failed": failed,
        "errors": errors[:10]
    }

async def import_price_file(file: UploadFile):
    try:
        contents = await file.read()
        filename = (file.filename or "").lower()

        if filename.endswith(('.xlsx', '.xls')):
            from openpyxl import load_workbook

            wb = load_workbook(BytesIO(contents))
            ws = wb.active
            rows = enumerate(ws.iter_rows(min_row=2, values_only=True), start=2)
            return import_price_rows(rows)

        if filename.endswith('.csv'):
            text = contents.decode('utf-8-sig')
            reader = csv.reader(text.splitlines())
            next(reader, None)
            rows = enumerate(reader, start=2)
            return import_price_rows(rows)

        raise HTTPException(status_code=400, detail="请上传 Excel 或 CSV 文件")
    except Exception as e:
        if isinstance(e, HTTPException):
            raise e
        raise HTTPException(status_code=500, detail=f"导入失败: {str(e)}")

@app.post("/api/import/prices")
async def import_prices(file: UploadFile = File(...)):
    return await import_price_file(file)

@app.post("/api/import/excel")
async def import_excel(file: UploadFile = File(...)):
    if not file.filename.endswith(('.xlsx', '.xls')):
        raise HTTPException(status_code=400, detail="请上传Excel文件 (.xlsx 或 .xls)")
    return await import_price_file(file)
