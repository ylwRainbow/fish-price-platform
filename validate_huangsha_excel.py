# -*- coding: utf-8 -*-
"""
黄颡鱼（汪丁鱼）历史价格数据验证脚本
验证Excel文件中数据的完整性和正确性
"""
import os
from datetime import datetime
from decimal import Decimal, InvalidOperation
from collections import defaultdict

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


EXCEL_PATH = "/Users/zcy/IdeaProjects/fish-price-platform/湖州黄颡鱼塘口价历史数据.xlsx"


def parse_price(price_str):
    """
    解析价格字符串
    
    Args:
        price_str: 价格字符串
        
    Returns:
        价格平均值或单个价格
    """
    if not price_str:
        return None
    
    price_str = str(price_str).strip()
    
    price_str = price_str.replace('元/斤', '').replace('元', '').replace('/斤', '')
    
    if '-' in price_str:
        parts = price_str.split('-')
        try:
            low = Decimal(parts[0].strip())
            high = Decimal(parts[1].strip())
            return (low + high) / 2
        except (InvalidOperation, ValueError, IndexError):
            return None
    
    try:
        return Decimal(price_str.strip())
    except (InvalidOperation, ValueError):
        return None


def validate_excel_data(file_path):
    """
    验证Excel数据
    
    Args:
        file_path: Excel文件路径
    """
    if load_workbook is None:
        print("错误: 需要安装openpyxl库")
        return
    
    print("=" * 60)
    print("黄颡鱼（汪丁鱼）历史价格数据验证")
    print("=" * 60)
    
    wb = load_workbook(file_path)
    ws = wb.active
    
    headers = []
    data = []
    
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
        if row_idx == 1:
            headers = [str(h) if h else f"col_{i}" for i, h in enumerate(row)]
            continue
        
        if not row[0]:
            continue
        
        row_data = {}
        for col_idx, value in enumerate(row):
            if col_idx < len(headers):
                row_data[headers[col_idx]] = value
        data.append(row_data)
    
    print(f"\n1. 数据概览")
    print(f"   总记录数: {len(data)} 条")
    
    print(f"\n2. 数据完整性检查")
    
    missing_date = 0
    missing_region = 0
    missing_price = 0
    missing_source = 0
    
    for row in data:
        if not row.get('日期'):
            missing_date += 1
        if not row.get('地区'):
            missing_region += 1
        if not row.get('塘口价(元/斤)'):
            missing_price += 1
        if not row.get('来源'):
            missing_source += 1
    
    print(f"   日期缺失: {missing_date} 条")
    print(f"   地区缺失: {missing_region} 条")
    print(f"   价格缺失: {missing_price} 条")
    print(f"   来源缺失: {missing_source} 条")
    
    print(f"\n3. 价格数据验证")
    
    invalid_prices = []
    abnormal_prices = []
    
    for row in data:
        price_str = row.get('塘口价(元/斤)', '')
        price = parse_price(price_str)
        
        if price is None:
            invalid_prices.append({
                'date': row.get('日期'),
                'region': row.get('地区'),
                'price': price_str
            })
        elif price > 30:
            abnormal_prices.append({
                'date': row.get('日期'),
                'region': row.get('地区'),
                'price': price_str,
                'parsed': float(price)
            })
    
    if invalid_prices:
        print(f"   无效价格: {len(invalid_prices)} 条")
        for p in invalid_prices[:5]:
            print(f"     - {p['date']} {p['region']}: {p['price']}")
    else:
        print(f"   无效价格: 0 条")
    
    if abnormal_prices:
        print(f"   异常高价(>30元/斤): {len(abnormal_prices)} 条")
        for p in abnormal_prices[:5]:
            print(f"     - {p['date']} {p['region']}: {p['price']}")
    else:
        print(f"   异常高价(>30元/斤): 0 条")
    
    print(f"\n4. 地区分布统计")
    
    region_stats = defaultdict(int)
    for row in data:
        region = row.get('地区', '未知')
        region_stats[region] += 1
    
    for region, count in sorted(region_stats.items(), key=lambda x: -x[1]):
        print(f"   {region}: {count} 条")
    
    print(f"\n5. 时间范围统计")
    
    years = defaultdict(int)
    for row in data:
        date_str = row.get('日期', '')
        if date_str:
            year = str(date_str)[:4]
            years[year] += 1
    
    for year, count in sorted(years.items()):
        print(f"   {year}年: {count} 条")
    
    print(f"\n6. 来源统计")
    
    source_stats = defaultdict(int)
    for row in data:
        source = row.get('来源', '未知')
        source_stats[source] += 1
    
    for source, count in sorted(source_stats.items(), key=lambda x: -x[1]):
        print(f"   {source}: {count} 条")
    
    print(f"\n7. 重复数据检查")
    
    seen = set()
    duplicates = []
    for row in data:
        key = (row.get('日期'), row.get('地区'), row.get('规格'), row.get('塘口价(元/斤)'))
        if key in seen:
            duplicates.append(row)
        else:
            seen.add(key)
    
    if duplicates:
        print(f"   发现重复数据: {len(duplicates)} 条")
        for d in duplicates[:5]:
            print(f"     - {d.get('日期')} {d.get('地区')} {d.get('规格')}: {d.get('塘口价(元/斤)')}")
    else:
        print(f"   重复数据: 0 条")
    
    print(f"\n8. 数据质量评估")
    
    total = len(data)
    valid_count = total - len(invalid_prices) - len(abnormal_prices)
    
    quality_score = (valid_count / total * 100) if total > 0 else 0
    
    if quality_score >= 95:
        quality_level = "优秀"
    elif quality_score >= 85:
        quality_level = "良好"
    elif quality_score >= 70:
        quality_level = "一般"
    else:
        quality_level = "较差"
    
    print(f"   数据质量评分: {quality_score:.1f}%")
    print(f"   数据质量等级: {quality_level}")
    
    print(f"\n" + "=" * 60)
    print("验证完成!")
    print("=" * 60)


if __name__ == "__main__":
    validate_excel_data(EXCEL_PATH)
